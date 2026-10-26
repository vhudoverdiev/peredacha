"""One-time retirement of the remark point, never of the worker category.

Only exact text within the same project/premise is deduplicated. Completion
is preserved; unmatched records are archived without guessing a work point.
"""
from collections import defaultdict
import json
import re

from sqlalchemy import text
from sqlalchemy.orm import selectinload

from app import db
from app.models import (AppSetting, ChangeLog, RemarkLibraryEntry,
                        RemarkPointReview, SyncConflict, SyncLog, Task, WorkPoint,
                        )
from app.services.remark_catalog import point_keys
from app.services.remark_text_cleaning import strip_worker_marker
from app.time_utils import utc_now
from app.work_points import WORK_POINT_LABELS, is_retired_point_label

MARKER = 'retire_remark_point_handyman_v1'


def text_key(value):
    # Unlike catalogue search, preserve all leading numbers and location words.
    return ' '.join(re.findall(r'[а-яa-z0-9]+', strip_worker_marker(value or '').casefold().replace('ё', 'е')))


def _same_state(a, b):
    # Two independent glass measurements must not overwrite one another.
    return not (a.glass_measurement and b.glass_measurement)


def _propagate_completion(source, target):
    if (source.status == 'done' or source.is_done) and not (target.status == 'done' and target.is_done):
        before = {'status': target.status, 'is_done': target.is_done,
                  'completed_date': str(target.completed_date or '')}
        target.status = 'done'
        target.is_done = True
        target.completed_date = source.completed_date or target.completed_date or utc_now()
        target.manually_edited = True
        db.session.add(ChangeLog(task=target, action='retired_point_completion_preserved',
                                 old_value=json.dumps(before), new_value=str(source.id)))


def _merge(duplicate, keeper):
    _propagate_completion(duplicate, keeper)
    db.session.add(ChangeLog(task=keeper, action='retired_point_state_preserved',
                             old_value=json.dumps({column.name: str(getattr(duplicate, column.name))
                                                   for column in Task.__table__.columns}, ensure_ascii=False)))
    for comment in list(duplicate.comments):
        comment.task = keeper
    for change in list(duplicate.changes):
        change.task = keeper
    if duplicate.comment and duplicate.comment != keeper.comment:
        keeper.comment = '\n\n'.join(filter(None, (keeper.comment, duplicate.comment)))
    if duplicate.glass_measurement:
        duplicate.glass_measurement.task = keeper
    # Do not lose material, source-conflict or repeat-measurement references.
    for writeoff in list(duplicate.material_writeoffs):
        if keeper not in writeoff.tasks:
            writeoff.tasks.append(keeper)
        writeoff.tasks.remove(duplicate)
    SyncConflict.query.filter_by(task_id=duplicate.id).update({'task_id': keeper.id}, synchronize_session=False)
    Task.query.filter_by(glass_parent_task_id=duplicate.id).update({'glass_parent_task_id': keeper.id}, synchronize_session=False)
    review = duplicate.remark_point_review
    if review:
        if keeper.remark_point_review is None:
            review.task = keeper
        else:
            db.session.add(ChangeLog(task=keeper, action='retired_point_review_preserved',
                                     old_value=json.dumps({'task_id':duplicate.id, 'action':review.action,
                                                          'fingerprint':review.fingerprint}, ensure_ascii=False)))
    db.session.add(ChangeLog(task=keeper, action='retired_point_duplicate_merged',
                             old_value=str(duplicate.id), new_value=str(keeper.id)))
    # Flush reassignment before delete-orphan cascades inspect the old task.
    db.session.flush()
    db.session.expire(duplicate, ['comments', 'changes', 'glass_measurement', 'remark_point_review'])
    db.session.delete(duplicate)
    db.session.flush()


def repair_points_from_source_headers():
    """Repair legacy physical-number mappings using the actual source cell only."""
    marker = AppSetting.query.filter_by(key=MARKER).first()
    if marker and json.loads(marker.value).get('policy_version') == 4:
        return 0
    from pathlib import Path
    from openpyxl import load_workbook
    from app.services.task_service import get_or_update_work_point
    repaired = 0
    projects = {row[0] for row in db.session.query(Task.project_id).distinct()}
    for project_id in projects:
        setting = AppSetting.query.filter_by(key=f'latest_excel_path_project_{project_id}').first()
        if not setting or not Path(setting.value).is_file():
            continue
        workbook = load_workbook(setting.value, read_only=True, data_only=True)
        try:
            tasks = Task.query.filter_by(project_id=project_id).all()
            for sheet in workbook:
                rows = list(sheet.iter_rows(values_only=True))
                header_options = []
                for row in rows[:30]:
                    mapping = {}
                    retired_columns = set()
                    for index, value in enumerate(row, 1):
                        label = str(value or '')
                        if is_retired_point_label(label):
                            retired_columns.add(index)
                            continue
                        keys = point_keys(label)
                        if len(keys) == 1 and keys[0] in WORK_POINT_LABELS:
                            mapping[index] = keys[0]
                    header_options.append((mapping, retired_columns))
                headers, retired_columns = max(header_options, key=lambda item: len(item[0]), default=({}, set()))
                if len(headers) < 3:
                    continue
                for task in tasks:
                    if task.source_sheet_name != sheet.title or not task.source_row_index or not task.source_column_index:
                        continue
                    column = task.source_column_index
                    number = headers.get(column)
                    if (not number and column not in retired_columns) or task.source_row_index > len(rows):
                        continue
                    row = rows[task.source_row_index - 1]
                    value = row[column - 1] if column <= len(row) else None
                    if not text_key(value) or text_key(value) != text_key(task.source_cell_value):
                        continue
                    if column in retired_columns:
                        # The source explicitly says Handyman: exclude this row,
                        # even if an older physical-number mapping labelled it otherwise.
                        retired = WorkPoint.query.filter_by(point_number='16', source_sheet_name='retired_source').first()
                        if retired is None:
                            retired = WorkPoint(point_number='16', short_name='Разнорабочие',
                                                original_column_name='Разнорабочие', source_sheet_name='retired_source', is_active=False)
                            db.session.add(retired)
                            db.session.flush()
                        old_id = task.work_point_id
                        task.work_point = retired
                        task.is_archived = True
                        db.session.add(ChangeLog(task=task, action='retired_source_column_excluded',
                                                 field_name='work_point_id', old_value=str(old_id), new_value=str(retired.id)))
                        repaired += 1
                        continue
                    # A deliberate manual correction remains authoritative.
                    if task.manually_edited:
                        continue
                    old = task.work_point
                    if old and old.point_number == number and point_keys(old.display_name) == [number]:
                        continue
                    target = get_or_update_work_point(number, WORK_POINT_LABELS[number], sheet.title, column)
                    task.work_point = target
                    if task.is_archived and ChangeLog.query.filter_by(task_id=task.id, action='retired_point_unmatched_archived').first():
                        task.is_archived = False
                    db.session.add(ChangeLog(task=task, action='source_header_point_repaired', field_name='work_point_id',
                                             old_value=str(old.id if old else ''), new_value=str(target.id)))
                    repaired += 1
                db.session.flush()
        finally:
            workbook.close()
    return repaired


def repair_legacy_named_points():
    """Some old imports used column 16 for PVC or another named work type.

    Retire by meaning, never discard those legitimate remarks by their old key.
    Keep IDs/source identities and union mapping relationships as on reimport.
    """
    mapping = {}
    for legacy in WorkPoint.query.filter_by(point_number='16').all():
        label = legacy.short_name or legacy.original_column_name or ''
        if is_retired_point_label(label):
            continue
        keys = point_keys(label)
        if len(keys) != 1 or keys[0] not in WORK_POINT_LABELS:
            continue
        number = keys[0]
        target = WorkPoint.query.filter_by(point_number=number).order_by(WorkPoint.id).first()
        if target is None:
            target = WorkPoint(point_number=number, short_name=WORK_POINT_LABELS[number],
                               original_column_name=WORK_POINT_LABELS[number], source_sheet_name='manual')
            db.session.add(target); db.session.flush()
        target.is_active = True
        for task in Task.query.filter_by(work_point_id=legacy.id).all():
            task.work_point = target
            db.session.add(ChangeLog(task=task, action='legacy_named_point_repaired', field_name='work_point_id',
                                     old_value=str(legacy.id),new_value=str(target.id)))
        for category in list(legacy.categories):
            if target not in category.work_points:
                category.work_points.append(target)
        for contractor in list(legacy.contractors):
            if target not in contractor.work_points:
                contractor.work_points.append(target)
        RemarkPointReview.query.filter_by(corrected_point_id=legacy.id).update(
            {'corrected_point_id':target.id}, synchronize_session=False)
        legacy.is_active = False
        mapping[str(legacy.id)] = number
    marker = AppSetting.query.filter_by(key=MARKER).first()
    if marker and mapping:
        summary = json.loads(marker.value)
        summary.setdefault('legacy_point_map', {}).update(mapping)
        marker.value = json.dumps(summary)
    db.session.info['retired_legacy_point_map'] = mapping
    db.session.flush()


def _run():
    marker = AppSetting.query.filter_by(key=MARKER).first()
    previous = json.loads(marker.value) if marker else {}
    if previous.get('policy_version') == 4:
        return
    points = WorkPoint.query.all()
    retired = [point for point in points if is_retired_point_label(point.short_name)
               or is_retired_point_label(point.original_column_name)
               or (point.point_number == '16' and not str(point.short_name or point.original_column_name or '').strip())]
    retired_ids = {point.id for point in retired}
    affected = Task.query.options(selectinload(Task.apartment), selectinload(Task.work_point)).filter(
        Task.work_point_id.in_(retired_ids)).order_by(Task.id).all() if retired_ids else []
    # Also reconcile rows transferred by the previous policy, without guessing a point.
    transferred_ids = {row.task_id for row in ChangeLog.query.filter_by(action='retired_point_relocated')}
    affected += Task.query.filter(Task.id.in_(transferred_ids), Task.work_point_id.notin_(retired_ids)).all() if transferred_ids else []
    affected_ids = {task.id for task in affected}
    from app.services.task_service import _premise_group_key
    def identity(task):
        if task.apartment:
            premise = (_premise_group_key(task.apartment),
                       str(task.apartment.building or '').strip(),
                       str(task.apartment.construction_number or '').strip())
        else:
            premise = f'without-premise:{task.id}'
        return (task.project_id, premise, text_key(task.description or task.source_cell_value))
    candidates = defaultdict(list)
    projects = {task.project_id for task in affected}
    for task in Task.query.options(selectinload(Task.apartment), selectinload(Task.work_point)).filter(
        Task.project_id.in_(projects), Task.work_point_id.notin_(retired_ids)).order_by(Task.id):
        if task.id not in affected_ids and task.work_point and task.work_point.point_number in WORK_POINT_LABELS:
            candidates[identity(task)].append(task)
    # Store the existing project snapshot inside the same transaction as changes.
    from app.services.sync_rollback import build_project_rollback_data
    for project_id in sorted(projects):
        log = SyncLog(project_id=project_id, source_type='maintenance', source_name=MARKER,
                      status='success', started_at=utc_now(), finished_at=utc_now(),
                      rollback_data=build_project_rollback_data(project_id),
                      rollback_note='Снимок перед переносом пунктов; восстановление разработчиком')
        snapshot = json.loads(log.rollback_data)
        snapshot['retired_work_points'] = [
            {column.name:getattr(point, column.name) for column in WorkPoint.__table__.columns
             if column.name not in {'created_at','updated_at'}} | {
                 'category_ids':[category.id for category in point.categories],
                 'contractor_ids':[contractor.id for contractor in point.contractors],
             } for point in retired
        ]
        log.rollback_data = json.dumps(snapshot, ensure_ascii=True)
        db.session.add(log)
    merged = quarantined = completed = 0
    for task in affected:
        key = identity(task)
        matches = candidates[key] if key[2] else []
        for other in matches:
            before = (other.status, other.is_done)
            _propagate_completion(task, other)
            completed += before != (other.status, other.is_done)
        keeper = next((other for other in matches if _same_state(task, other)), None)
        if keeper:
            _merge(task, keeper)
            merged += 1
            continue
        # Unexpected unmatched data is retained for investigation, never classified
        # as Misc or silently deleted. The original source and all links survive.
        task.is_archived = True
        db.session.add(ChangeLog(task=task, action='retired_point_unmatched_archived',
                                 old_value=str(task.work_point_id)))
        quarantined += 1
    # Remove catalogue examples belonging to the retired point, rather than
    # silently teaching those phrases as another work type.
    for entry in RemarkLibraryEntry.query.all():
        if '16' in json.loads(entry.points_json):
            db.session.delete(entry)
    db.session.flush()
    for point in retired:
        RemarkPointReview.query.filter_by(corrected_point_id=point.id).update(
            {'corrected_point_id':None}, synchronize_session=False)
        point.categories.clear(); point.contractors.clear()
        db.session.flush()
        if Task.query.filter_by(work_point_id=point.id).first():
            point.is_active = False
        else:
            db.session.delete(point)
    summary = {'policy_version':4, 'moved':0, 'merged':merged, 'quarantined':quarantined,
               'completed':completed, 'retired_points':len(retired),
               'retired_point_ids':sorted(retired_ids | set(previous.get('retired_point_ids', []))),
               'legacy_point_map':previous.get('legacy_point_map', {}) | db.session.info.get('retired_legacy_point_map', {})}
    if marker:
        marker.value = json.dumps(summary)
    else:
        db.session.add(AppSetting(key=MARKER, value=json.dumps(summary)))
    db.session.commit()
    return summary


def retire_handyman_remark_point():
    # Gunicorn workers must not migrate the same tasks concurrently.
    db.session.rollback()
    db.session.info['retiring_work_point'] = True
    if db.engine.dialect.name == 'mysql':
        with db.engine.connect() as lock:
            if lock.execute(text("SELECT GET_LOCK(:key, 120)"), {'key':MARKER}).scalar() != 1:
                raise RuntimeError('Не удалось заблокировать перенос замечаний для обслуживания.')
            try:
                repair_points_from_source_headers()
                repair_legacy_named_points()
                result = _run()
                purge_retired_library_examples()
                db.session.commit()
                return result
            except Exception:
                db.session.rollback(); raise
            finally:
                db.session.info.pop('retiring_work_point', None)
                db.session.info.pop('retired_legacy_point_map', None)
                lock.execute(text('SELECT RELEASE_LOCK(:key)'), {'key':MARKER})
    else:
        try:
            repair_points_from_source_headers()
            repair_legacy_named_points()
            result = _run()
            purge_retired_library_examples()
            db.session.commit()
            return result
        except Exception:
            db.session.rollback(); raise
        finally:
            db.session.info.pop('retiring_work_point', None)
            db.session.info.pop('retired_legacy_point_map', None)


def purge_retired_library_examples():
    """Also clean installations which already ran the earlier retirement."""
    key = 'retired_point_library_cleanup_v2'
    if AppSetting.query.filter_by(key=key).first():
        return
    from app.work_points import retired_phrase_hashes
    hashes = retired_phrase_hashes()
    removed = 0
    for entry in RemarkLibraryEntry.query.all():
        if '16' in json.loads(entry.points_json) or entry.normalized_hash in hashes:
            db.session.delete(entry); removed += 1
    db.session.add(AppSetting(key=key, value=str(removed)))
    db.session.commit()
