"""Rebuild the shipped reference from explicitly supplied remark columns only.

Usage: python -m app.services.remark_catalog_import INPUT.xlsx ...
"""
from hashlib import sha256
import json
from pathlib import Path
import re
import sys

from openpyxl import load_workbook
from app.services.remark_catalog import CATALOG_PATH, normalize, point_keys
from app.services.uid_service import split_cell_remarks


def build_catalog(paths):
    if not paths:
        raise ValueError('Укажите исходные Excel-файлы. Существующий справочник не изменён.')
    entries={}; sources=[]; stats=[]
    for source_id,path in enumerate(paths):
        path=Path(path)
        sources.append({'name':path.name,'sha256':sha256(path.read_bytes()).hexdigest()})
        workbook=load_workbook(path,read_only=True,data_only=True)
        observations=0
        try:
            # Dedicated contractor sheets mirror the same remarks and can be
            # stale; learn from the primary registers, not their formula copies.
            for sheet in workbook:
                if sheet.title.casefold() not in {'таблица','коммерция','кладовые','паркинг'}: continue
                rows=list(sheet.iter_rows(values_only=True))
                if not rows: continue
                header=max(range(min(12,len(rows))),key=lambda i:sum(bool(point_keys(v)) for v in rows[i]))
                columns={i:point_keys(v) for i,v in enumerate(rows[header]) if point_keys(v) and '26' not in point_keys(v)}
                for row_index,row in enumerate(rows[header+1:],start=header+2):
                    for col,points in columns.items():
                        value=row[col] if col<len(row) else None
                        for fragment in split_cell_remarks(value):
                            text=re.sub(r'\s+', ' ',fragment).strip(' -–—\"«»')
                            text=re.sub(r'\b[\w.+-]+@[\w.-]+\.[a-z]{2,}\b','[контакт удалён]',text,flags=re.I)
                            text=re.sub(r'(?<!\d)(?:\+7|8|7)[\s(-]*\d{3}[)\s-]*\d{3}[\s-]*\d{2}[\s-]*\d{2}(?!\d)','[контакт удалён]',text)
                            key=normalize(text)
                            if len(key)<12 or len(text)>1200 or not re.search('[а-яА-Я]',text): continue
                            item=entries.setdefault(key,{'text':text,'points':set(),'count':0,'origins':[]})
                            item['points'].update(points); item['count']+=1; observations+=1
                            if len(item['origins'])<4:
                                raw_number=rows[header+1][col] if header+1<len(rows) else None
                                item['origins'].append({'source':source_id,'sheet':sheet.title,'row':row_index,'column':col+1,
                                                        'source_number':str(raw_number or ''), 'label':str(rows[header][col]).strip()})
        finally: workbook.close()
        stats.append({'file':path.name,'observations':observations})
    from app.services.remark_text_cleaning import clean_reference, vocabulary
    words=vocabulary(item['text'] for item in entries.values())
    cleaned={}
    for item in entries.values():
        item['text']=clean_reference(item['text'],words)
        key=normalize(item['text'])
        if key in cleaned:
            saved=cleaned[key];saved['points'].update(item['points']);saved['count']+=item['count']
            saved['origins']=(saved['origins']+item['origins'])[:4]
        else: cleaned[key]=item
    entries=cleaned
    result=[]
    for key,item in sorted(entries.items()):
        item['points']=sorted(item['points'],key=int)
        item['id']=sha256(key.encode()).hexdigest()[:16]
        result.append(item)
    version=sha256(json.dumps({'sources':sources,'entries':result},sort_keys=True,ensure_ascii=False).encode()).hexdigest()[:16]
    if not result:
        raise ValueError('В указанных файлах не найдены замечания в именованных колонках. Существующий справочник не изменён.')
    return {'version':version,'sources':sources,'entries':result},stats


if __name__=='__main__':
    payload,stats=build_catalog(sys.argv[1:])
    CATALOG_PATH.parent.mkdir(parents=True,exist_ok=True)
    header=json.dumps({'version':payload['version'],'sources':payload['sources']},ensure_ascii=False)[:-1]
    records=',\n'.join(json.dumps(entry,ensure_ascii=False,separators=(',',':')) for entry in payload['entries'])
    CATALOG_PATH.write_text(header+',"entries":[\n'+records+'\n]}\n',encoding='utf-8')
    print(json.dumps({'sources':stats,'unique':len(payload['entries']),'ambiguous':sum(len(e['points'])>1 for e in payload['entries'])},ensure_ascii=False,indent=2))
