from io import BytesIO
from pathlib import Path
from unittest.mock import patch
import unittest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from app import db
from app.models import STATUS_CONCESSION, STATUS_FINISHERS, STATUS_GUARANTEE, STATUS_NOT_STARTED, Task, WorkPoint
import test_worker_task_updates as fixture


class ApartmentSourceExportTests(unittest.TestCase):
    setUp = fixture.WorkerTaskUpdatesTests.setUp
    tearDown = fixture.WorkerTaskUpdatesTests.tearDown
    login = fixture.WorkerTaskUpdatesTests.login

    def export_from_headers(self, headers, prefix_rows=(), index_numbers=None):
        from app.services.excel_export import export_apartment_source_table
        source=Path(self.temp_files.name)/'changed-template.xlsx'
        workbook=Workbook(); sheet=workbook.active
        for row in prefix_rows:
            sheet.append(row)
        sheet.append(headers)
        if index_numbers:
            sheet.append(index_numbers)
        sheet.append(['999', 'Данные другого помещения'])
        workbook.save(source); workbook.close()
        with patch('app.services.excel_export.resolve_source_excel_with_strikes_path', return_value=source):
            path=export_apartment_source_table([self.apartment], [self.task], 'name-mapping')
        self.addCleanup(path.unlink, missing_ok=True)
        exported=load_workbook(path, rich_text=True)
        self.addCleanup(exported.close)
        return exported.active

    def test_changed_numbers_and_column_order_match_by_name(self):
        self.point.point_number='11'
        self.point.short_name='Электрика'
        sheet=self.export_from_headers(['№ квартиры', '11. Вентиляция', '24. Электрика'])
        self.assertFalse(sheet['B2'].value)
        self.assertIn('Покрасить стену', str(sheet['C2'].value))

    def test_names_ignore_number_case_punctuation_and_yo(self):
        self.point.short_name='11. Шпаклёвка стен'
        sheet=self.export_from_headers(['№ квартиры', '24. ШПАКЛЕВКА\nСТЕН'])
        self.assertIn('Покрасить стену', str(sheet['B2'].value))
        self.assertEqual(sheet.max_column, 2)

    def test_missing_named_column_is_added_instead_of_using_same_number(self):
        self.point.point_number='11'
        self.point.short_name='Электрика'
        sheet=self.export_from_headers(['№ квартиры', '11. Вентиляция'])
        self.assertFalse(sheet['B2'].value)
        self.assertEqual(sheet['C1'].value, 'Электрика')
        self.assertIn('Покрасить стену', str(sheet['C2'].value))

    def test_summary_rows_and_columns_are_not_copied(self):
        self.point.short_name='Электрика'
        sheet=self.export_from_headers(
            ['№ квартиры', '24. Электрика', '27. Всего помещений', '28. Передано', '29. Не передано'],
            [['Всего помещений', 900], ['Передано', 800], ['Не передано', 100]],
        )
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.max_column, 2)
        self.assertEqual(sheet['A1'].value, '№ квартиры')
        self.assertIn('Покрасить стену', str(sheet['B2'].value))
        self.assertNotIn('передано', str(list(sheet.values)).lower())

    def test_imported_full_name_is_used_when_short_label_differs(self):
        self.point.short_name='Витражи'
        self.point.original_column_name='20. Холодное остекление балконов'
        sheet=self.export_from_headers(['№ квартиры', '17. Холодное остекление балконов'])
        self.assertEqual(sheet.max_column, 2)
        self.assertIn('Покрасить стену', str(sheet['B2'].value))

    def test_base_finishing_column_is_not_mistaken_for_a_remark_point(self):
        self.point.short_name='Отделка'
        sheet=self.export_from_headers(['№ квартиры', 'Отделка', '11. Отделка'])
        self.assertNotIn('Покрасить стену', str(sheet['B2'].value))
        self.assertIn('Покрасить стену', str(sheet['C2'].value))

    def test_duplicate_names_are_not_assigned_arbitrarily(self):
        self.point.short_name='Электрика'
        sheet=self.export_from_headers(['№ квартиры', '11. Электрика', '24. Электрика'])
        self.assertFalse(sheet['B2'].value)
        self.assertFalse(sheet['C2'].value)
        self.assertEqual(sheet['D1'].value, 'Электрика')
        self.assertIn('Покрасить стену', str(sheet['D2'].value))

    def test_number_row_does_not_rewrite_labels_from_another_template(self):
        self.point.point_number='21'
        self.point.short_name='Электрика'
        sheet=self.export_from_headers(
            ['№ квартиры', 'Система отопления', 'Канализация', 'Прочее', 'Вентиляция', 'Электрика'],
            index_numbers=[1, 21, 22, 23, 24, 25],
        )
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet['F1'].value, 'Электрика')
        self.assertEqual(list(sheet.values)[1], (1,20,21,24,10,23))
        self.assertFalse(sheet['B3'].value)
        self.assertIn('Покрасить стену', str(sheet['F3'].value))

    def test_merged_header_fill_numbers_and_missing_points_stay_before_register(self):
        from app.services.excel_export import export_apartment_source_table
        from app.work_points import WORK_POINT_LABELS
        from openpyxl.styles import Alignment
        source=Path(self.temp_files.name)/'merged-register.xlsx'
        wb=Workbook(); ws=wb.active
        ws.append(['Всего помещений', 900, 'Передано', 800])
        ws.append(['№ квартиры', 'Собственник', 'Номер телефона', 'Вид отделки'])
        ws.append([None,None,None,None,WORK_POINT_LABELS['13'],WORK_POINT_LABELS['15'],
                   WORK_POINT_LABELS['18'],WORK_POINT_LABELS['20'],WORK_POINT_LABELS['25'],
                   'Дата занесения в РВТ (ПТО)'])
        for col in range(1,5):
            ws.merge_cells(start_row=2,end_row=3,start_column=col,end_column=col)
        for row in (2,3):
            for cell in ws[row]:
                if cell.value:
                    cell.fill=PatternFill('solid',fgColor='93C47D')
                    cell.font=Font(bold=True)
                    cell.alignment=Alignment(wrap_text=True)
        ws['J3'].fill=PatternFill('solid',fgColor='FBBC04')
        ws.append(list(range(1,11)))
        ws.append(['999','Не экспортировать'])
        wb.save(source); wb.close()
        self.point.short_name=WORK_POINT_LABELS['14']
        self.point.point_number='99'
        other=WorkPoint(point_number='98',short_name=WORK_POINT_LABELS['19'],source_sheet_name='manual')
        task=Task(source_uid='export-missing-rostelecom',project=self.project,apartment=self.apartment,
                  work_point=other,description='Подключить интернет',status=STATUS_NOT_STARTED)
        db.session.add(task); db.session.commit()
        with patch('app.services.excel_export.resolve_source_excel_with_strikes_path',return_value=source):
            path=export_apartment_source_table([self.apartment],[self.task,task],'merged-regression')
        self.addCleanup(path.unlink,missing_ok=True)
        result=load_workbook(path); self.addCleanup(result.close)
        sheet=result.active
        headers=[c.value for c in sheet[1]]
        self.assertEqual(headers[4:11], [WORK_POINT_LABELS[n] for n in ('13','14','15','18','19','20','25')])
        self.assertEqual(headers[-1],'Дата занесения в РВТ (ПТО)')
        self.assertEqual(sheet.max_row,3)
        self.assertEqual([c.value for c in sheet[2]],[1,2,3,4,13,14,15,17,18,19,24,12])
        for col in range(1,5):
            self.assertEqual(sheet.cell(1,col).fill.fgColor.rgb,'0093C47D')
            self.assertTrue(sheet.cell(1,col).font.bold)
        self.assertEqual(sheet.cell(1,12).fill.fgColor.rgb,'00FBBC04')
        self.assertIn('Покрасить стену',str(sheet.cell(3,6).value))
        self.assertIn('Подключить интернет',str(sheet.cell(3,9).value))

    def test_repeat_exports_have_distinct_paths_and_keep_the_first_workbook_intact(self):
        from app.services.excel_export import export_apartment_source_table
        with patch('app.services.excel_export.resolve_source_excel_with_strikes_path', side_effect=FileNotFoundError):
            first = export_apartment_source_table([self.apartment], [self.task], 'same-premise')
            original = first.read_bytes()
            self.task.description = 'Новое описание'
            second = export_apartment_source_table([self.apartment], [self.task], 'same-premise')
        self.assertNotEqual(first, second)
        self.assertEqual(first.read_bytes(), original)
        for path in [first, second]:
            workbook = load_workbook(path)
            self.assertEqual(workbook.active.max_row, 2)
            workbook.close()

    def test_long_cyrillic_project_name_does_not_exceed_filesystem_limit(self):
        from app.services.excel_export import export_apartment_source_table
        with patch('app.services.excel_export.resolve_source_excel_with_strikes_path', side_effect=FileNotFoundError):
            path = export_apartment_source_table([self.apartment], [self.task], 'Очень длинный объект ' * 20)
        self.assertLess(len(path.name.encode('utf-8')), 255)
        workbook = load_workbook(path)
        self.assertEqual(workbook.active.max_row, 2)
        workbook.close()

    def test_template_headers_single_current_row_and_styles(self):
        self.login(self.manager)
        source=Path(self.temp_files.name)/'source.xlsx'
        wb=Workbook();ws=wb.active;ws.title='Квартиры'
        ws.append(['№ квартиры','ФИО собственника','Телефон','11. Стены и потолки'])
        ws['A1'].font=Font(name='Arial',bold=True,color='123456')
        ws['A1'].fill=PatternFill('solid',fgColor='AABBCC')
        ws['AR1'].fill=PatternFill('solid',fgColor='FFC000')
        ws.auto_filter.ref='A1:AR2'
        ws.append(['999','Чужой собственник','other phone','Секретное замечание'])
        hidden=wb.create_sheet('Hidden');hidden.sheet_state='hidden';hidden.append(['Other secret'])
        wb.save(source);wb.close()
        self.point.point_number='11'; self.point.short_name='Стены и потолки'
        self.task.is_done=True;self.task.status='done';db.session.commit()
        with patch('app.services.excel_export.resolve_source_excel_with_strikes_path',return_value=source):
            response=self.client.get(f'/apartments/{self.apartment.id}/source-table/export')
        self.assertEqual(response.status_code,200)
        exported=load_workbook(BytesIO(response.data),rich_text=True)
        self.assertEqual(exported.sheetnames,['Квартиры'])
        sheet=exported.active
        self.assertEqual(sheet.max_row,2)
        self.assertEqual(sheet.max_column,4)
        self.assertEqual(sheet.auto_filter.ref,'A1:D2')
        self.assertEqual(sheet.print_area,"'Квартиры'!$A$1:$D$2")
        self.assertEqual(sheet['A2'].value,'173')
        self.assertTrue(sheet['A1'].font.bold)
        self.assertEqual(sheet['A1'].fill.fgColor.rgb,'00AABBCC')
        content=' '.join(str(c.value) for row in sheet for c in row)
        self.assertIn('Покрасить стену',content)
        self.assertIn('- Покрасить стену', content)
        self.assertNotIn('Секретное',content);self.assertNotIn('999',content)
        self.assertTrue(any(c.font.strike for c in sheet[2]))
        exported.close();response.close()

    def test_one_apartment_status_markers_match_source_table_rules(self):
        self.login(self.manager)
        self.point.point_number = '11'
        self.task.description = 'Устранённое замечание'
        self.task.status = 'done'
        self.task.is_done = True
        statuses = [
            ('finishers', STATUS_FINISHERS, 'Передать отделочникам'),
            ('concession', STATUS_CONCESSION, 'Согласованное отступление'),
            ('guarantee', STATUS_GUARANTEE, 'Гарантийная работа'),
            ('open', STATUS_NOT_STARTED, 'Открытое замечание'),
        ]
        for uid, status, description in statuses:
            db.session.add(Task(
                source_uid=f'apartment-export-{uid}', project=self.project,
                apartment=self.apartment, work_point=self.point,
                description=description, source_cell_value=description,
                status=status, is_done=status != STATUS_NOT_STARTED,
            ))
        db.session.commit()
        response = self.client.get(f'/apartments/{self.apartment.id}/source-table/export')
        workbook = load_workbook(BytesIO(response.data), rich_text=True)
        text = '\n'.join(str(cell.value or '') for cell in workbook.active[workbook.active.max_row])
        self.assertIn('- Устранённое замечание', text)
        self.assertIn('Передать отделочникам (чистовики)', text)
        self.assertIn('Согласованное отступление (отступное)', text)
        self.assertIn('Гарантийная работа (гарантия)', text)
        self.assertIn('Открытое замечание', text)
        self.assertNotIn('- Открытое замечание', text)
        workbook.close(); response.close()

    def test_fallback_without_source_and_worker_denial(self):
        self.login(self.manager)
        page = self.client.get(f'/apartments/{self.apartment.id}').get_data(as_text=True)
        self.assertIn('btn btn-success apartment-source-export-btn', page)
        self.assertNotIn('title="Скачать таблицу только этого помещения"', page)
        with patch('app.services.excel_export.resolve_source_excel_with_strikes_path',side_effect=FileNotFoundError):
            response=self.client.get(f'/apartments/{self.apartment.id}/source-table/export')
        self.assertEqual(response.status_code,200)
        workbook=load_workbook(BytesIO(response.data));self.assertEqual(workbook.active.max_row,2)
        self.assertIn('Покрасить стену',' '.join(str(c.value) for c in workbook.active[2]))
        workbook.close();response.close()
        self.login(self.worker)
        self.assertNotEqual(self.client.get(f'/apartments/{self.apartment.id}/source-table/export').status_code,200)
