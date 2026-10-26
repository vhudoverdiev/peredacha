import hashlib
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook
from sqlalchemy import text

from app import create_app, db
from app.models import (AppSetting, Apartment, ChangeLog, Project, RemarkLibraryEntry,
                        SyncConflict, Task, TaskComment, User, WorkCategory, WorkPoint)
from app.services.excel_export import export_apartment_source_table, export_source_excel_with_strikes
from app.services.excel_import import mark_missing_tasks, sync_excel_file
from app.services.mapping_service import ensure_default_categories
from app.services.remark_catalog import get_base_catalog
from app.services.remark_library_file import import_library, SCHEMA
from app.services.retired_work_point import MARKER, retire_handyman_remark_point
from app.services.task_service import set_setting, sync_rows
from app.work_points import WORK_POINT_LABELS, display_point_number
from config import Config


class RetirementTests(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        root = Path(self.directory.name)
        class TestConfig(Config):
            TESTING = True
            SECRET_KEY = 'retirement-test'
            SQLALCHEMY_DATABASE_URI = 'sqlite://'
            WTF_CSRF_ENABLED = False
            UPLOAD_FOLDER = str(root / 'uploads')
            EXPORT_FOLDER = str(root / 'exports')
        self.app = create_app(TestConfig)
        self.context = self.app.app_context(); self.context.push()
        db.session.execute(text('PRAGMA foreign_keys=ON'))
        AppSetting.query.filter_by(key=MARKER).delete()
        self.project = Project(name='Retirement QA')
        self.apartment = Apartment(project=self.project, apartment_number='42', construction_number='1-2-1')
        self.retired = WorkPoint(point_number='16', short_name='Разнорабочие', source_sheet_name='old')
        self.pvc = WorkPoint(point_number='17', short_name=WORK_POINT_LABELS['17'], source_sheet_name='old')
        self.user = User(username='retirement-user', role='admin', password_hash='unused')
        db.session.add_all([self.project,self.apartment,self.retired,self.pvc,self.user]); db.session.commit()

    def tearDown(self):
        db.session.remove(); db.drop_all(); self.context.pop(); self.directory.cleanup()

    def task(self, uid, point, description, **values):
        task = Task(source_uid=uid, project=self.project, apartment=self.apartment,
                    work_point=point, description=description, **values)
        db.session.add(task); db.session.flush(); return task

    def migrate(self):
        db.session.commit()
        return retire_handyman_remark_point()

    def source(self):
        workbook=Workbook(); sheet=workbook.active; sheet.title='Таблица'
        sheet.append(['№ кв. проект','стр. №','Собственник','Вид отделки','Вентиляция','Разнорабочие',WORK_POINT_LABELS['17'],'Электрика'])
        sheet.append([1,2,3,4,10,16,17,24])
        sheet.append(['42','1-2-1','Собственник','Белая','Нет вентиляции','Дубль из удалённого пункта','Регулировка ПВХ','Повреждена розетка'])
        path=Path(self.directory.name)/'source.xlsx';workbook.save(path)
        return path

    def assert_no_retired_header(self, workbook):
        self.assertFalse(any('разнорабоч' in str(cell.value or '').casefold()
                             for sheet in workbook for row in sheet for cell in row))

    def test_duplicates_merge_without_losing_comments_history_or_conflicts(self):
        keeper=self.task('keeper',self.pvc,'Регулировка ПВХ')
        duplicate=self.task('duplicate',self.retired,'Регулировка ПВХ')
        comment=TaskComment(task=duplicate,user=self.user,body='Сохранить комментарий')
        history=ChangeLog(task=duplicate,action='existing_history')
        conflict=SyncConflict(task=duplicate,old_value='old',new_value='new')
        db.session.add_all([comment,history,conflict]);db.session.commit()
        ids=(keeper.id,duplicate.id,comment.id,history.id,conflict.id)
        result=self.migrate(); db.session.expire_all()
        self.assertEqual(result['merged'],1)
        self.assertIsNone(db.session.get(Task,ids[1]))
        self.assertEqual(db.session.get(TaskComment,ids[2]).task_id,ids[0])
        self.assertEqual(db.session.get(ChangeLog,ids[3]).task_id,ids[0])
        self.assertEqual(db.session.get(SyncConflict,ids[4]).task_id,ids[0])
        self.assertIsNone(WorkPoint.query.filter_by(point_number='16').first())
        self.assertIsNone(retire_handyman_remark_point())

    def test_unmatched_tasks_are_archived_without_inferred_point(self):
        unique=self.task('unique',self.retired,'Повреждена рама ПВХ')
        unknown=self.task('unknown',self.retired,'Необычный дефект квазипокрытия')
        ids=(unique.id,unknown.id)
        self.migrate();db.session.expire_all()
        for task_id in ids:
            task = db.session.get(Task,task_id)
            self.assertEqual(task.work_point.point_number,'16')
            self.assertTrue(task.is_archived)
            self.assertFalse(task.work_point.is_active)
        self.assertIsNone(WorkPoint.query.filter_by(point_number='25').first())

    def test_distinct_operational_state_is_not_deleted_as_duplicate(self):
        self.task('done',self.pvc,'Регулировка ПВХ',status='done',is_done=True)
        self.task('open',self.retired,'Регулировка ПВХ')
        result=self.migrate()
        self.assertEqual(result['merged'],1)
        self.assertEqual(Task.query.count(),1)
        self.assertTrue(Task.query.one().is_done)

    def test_done_duplicate_completes_all_exact_matches_only(self):
        from app.time_utils import utc_now
        completed = utc_now()
        keeper = self.task('open-target', self.pvc, 'Регулировка ПВХ', status='new', is_done=False)
        another = WorkPoint(point_number='23', short_name=WORK_POINT_LABELS['23'])
        db.session.add(another)
        second = self.task('second-target', another, 'Регулировка ПВХ', is_done=False)
        unrelated = self.task('different-text', self.pvc, 'Другая регулировка ПВХ', is_done=False)
        self.task('done-source', self.retired, 'Регулировка ПВХ', status='done', is_done=True, completed_date=completed)
        ids = keeper.id, second.id, unrelated.id
        result = self.migrate()
        db.session.expire_all()
        self.assertEqual(result['merged'], 1)
        for task_id in ids[:2]:
            task = db.session.get(Task, task_id)
            self.assertEqual(task.status, 'done')
            self.assertTrue(task.is_done)
            self.assertEqual(task.completed_date.replace(tzinfo=None), completed.replace(tzinfo=None))
        self.assertFalse(db.session.get(Task, ids[2]).is_done)

    def test_previous_policy_transfer_is_reconciled_on_upgrade(self):
        keeper = self.task('upgrade-target', self.pvc, 'Регулировка ПВХ', is_done=False)
        old = self.task('upgrade-source', self.pvc, 'Регулировка ПВХ', status='done', is_done=True,
                        source_sheet_name='manual_retired_point')
        db.session.add(ChangeLog(task=old, action='retired_point_relocated'))
        db.session.add(AppSetting(key=MARKER, value=json.dumps({'retired_point_ids':[999], 'merged':0})))
        keeper_id = keeper.id
        result = self.migrate()
        self.assertEqual(result['merged'], 1)
        self.assertTrue(db.session.get(Task, keeper_id).is_done)
        self.assertIn(999, result['retired_point_ids'])
        self.assertIsNone(retire_handyman_remark_point())

    def test_source_header_restores_pvc_from_overwritten_shared_point(self):
        path = self.source()
        task = self.task('legacy-pvc-cell', self.retired, 'Регулировка ПВХ',
                         source_sheet_name='Таблица', source_row_index=3,
                         source_column_index=7, source_cell_value='Регулировка ПВХ', is_archived=True)
        db.session.add(ChangeLog(task=task, action='retired_point_unmatched_archived'))
        set_setting(f'latest_excel_path_project_{self.project.id}', str(path))
        task_id = task.id
        self.migrate()
        db.session.expire_all()
        restored = db.session.get(Task, task_id)
        self.assertEqual(restored.work_point.point_number, '17')
        self.assertFalse(restored.is_archived)
        self.assertEqual(restored.source_uid, 'legacy-pvc-cell')

    def test_retired_source_column_is_excluded_even_if_old_point_is_valid(self):
        path = self.source()
        task = self.task('old-shifted-handyman', self.pvc, 'Дубль из удалённого пункта',
                         source_sheet_name='Таблица', source_row_index=3,
                         source_column_index=6, source_cell_value='Дубль из удалённого пункта')
        set_setting(f'latest_excel_path_project_{self.project.id}', str(path))
        task_id = task.id
        self.migrate()
        db.session.expire_all()
        task = db.session.get(Task, task_id)
        self.assertTrue(task.is_archived)
        self.assertEqual(task.work_point.point_number, '16')
        self.assertFalse(task.work_point.is_active)
        self.assertIsNone(WorkPoint.query.filter_by(point_number='25').first())

    def test_library_examples_are_deleted_not_relabelled_and_cannot_reimport(self):
        for i,points in enumerate([['16'],['16','17'],['17']]):
            db.session.add(RemarkLibraryEntry(normalized_hash=str(i),point_key=','.join(points),
                text=f'Специальный пример {i}',points_json=json.dumps(points),occurrences=1))
        self.migrate()
        self.assertEqual([row.text for row in RemarkLibraryEntry.query.all()],['Специальный пример 2'])
        path=Path(self.directory.name)/'phrases.json'
        path.write_text(json.dumps({'schema':SCHEMA,'entries':[{'text':'Старый пример не возвращается','points':['16']}]}))
        import_library(path)
        self.assertEqual(RemarkLibraryEntry.query.count(),1)
        self.assertTrue(all('16' not in row['points'] for row in get_base_catalog().entries))

    def test_old_workbook_column_is_ignored_and_neighbouring_points_import(self):
        path=self.source()
        self.migrate()
        sync_excel_file(path,project_name=self.project.name)
        tasks=Task.query.all()
        self.assertEqual({task.description for task in tasks},{'Нет вентиляции','Регулировка ПВХ','Повреждена розетка'})
        self.assertEqual({task.work_point.point_number for task in tasks},{'10','17','24'})
        self.assertEqual(Apartment.query.filter_by(apartment_number='42').first().construction_number,'1-2-1')

    def test_new_contiguous_number_row_imports_by_name_not_number(self):
        self.migrate()
        sync_rows([
            ['№ кв. проект','стр. №','Собственник','Вид отделки','Вентиляция',WORK_POINT_LABELS['17'],'Электрика'],
            [1,2,3,4,10,16,23],
            ['42','1-2-1','Владелец','Белая','Нет вентиляции','Регулировка ПВХ','Повреждена розетка'],
        ],sheet_name='Таблица',project_name=self.project.name)
        self.assertEqual({task.work_point.point_number for task in Task.query.all()},{'10','17','24'})
        self.assertEqual([display_point_number(number) for number in WORK_POINT_LABELS if int(number)>=10],
                         [str(number) for number in range(10,26)])

    def test_both_exports_drop_point_and_keep_pvc_and_original_file(self):
        path=self.source(); original=hashlib.sha256(path.read_bytes()).hexdigest()
        self.task('pvc-export',self.pvc,'Регулировка ПВХ')
        self.task('unique-export',self.pvc,'Необычный дефект квазипокрытия', source_sheet_name='manual_retired_point')
        self.migrate()
        set_setting('latest_excel_path',str(path))
        set_setting(f'latest_excel_path_project_{self.project.id}',str(path));db.session.commit()
        for exported in [
            export_apartment_source_table([self.apartment],Task.query.all(),'apartment'),
            export_source_excel_with_strikes(str(path),self.project.name,self.project.id),
        ]:
            workbook=load_workbook(exported,data_only=True)
            self.assert_no_retired_header(workbook)
            values=[str(cell.value or '') for sheet in workbook for row in sheet for cell in row]
            self.assertTrue(any('Регулировка ПВХ' in value for value in values))
            self.assertTrue(any('Необычный дефект квазипокрытия' in value for value in values))
            workbook.close()
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),original)

    def test_worker_category_remains(self):
        self.migrate();ensure_default_categories();db.session.commit()
        self.assertTrue(WorkCategory.query.filter_by(name='Разнорабочие').one().is_active)


    def test_rollback_of_failed_migration_restores_tasks_and_point(self):
        from unittest.mock import patch
        old=self.task('rollback-unique',self.retired,'Контрольный неизвестный дефект')
        db.session.commit(); old_id=old.id; point_id=self.retired.id
        with patch('app.services.retired_work_point.text_key',side_effect=RuntimeError('simulated failure')):
            with self.assertRaises(RuntimeError):retire_handyman_remark_point()
        db.session.expire_all()
        self.assertEqual(db.session.get(Task,old_id).work_point_id,point_id)
        self.assertIsNotNone(db.session.get(WorkPoint,point_id))
        self.assertIsNone(AppSetting.query.filter_by(key=MARKER).first())

    def test_glass_material_and_review_links_survive_merging(self):
        from app.models import GlassMeasurement, MaterialWriteOff, RemarkPointReview
        keeper=self.task('linked-keeper',self.pvc,'Регулировка ПВХ')
        duplicate=self.task('linked-duplicate',self.retired,'Регулировка ПВХ')
        glass=GlassMeasurement(project=self.project,task=duplicate,apartment=self.apartment,width=500,height=700)
        writeoff=MaterialWriteOff(project=self.project,tasks=[duplicate])
        review=RemarkPointReview(task=duplicate,fingerprint='fingerprint',catalog_version='v1',action='corrected')
        db.session.add_all([glass,writeoff,review]);db.session.commit()
        ids=keeper.id,glass.id,writeoff.id,review.id
        self.migrate();db.session.expire_all()
        self.assertEqual(db.session.get(GlassMeasurement,ids[1]).task_id,ids[0])
        self.assertEqual([task.id for task in db.session.get(MaterialWriteOff,ids[2]).tasks],[ids[0]])
        self.assertEqual(db.session.get(RemarkPointReview,ids[3]).task_id,ids[0])

    def test_cleanup_handles_examples_relabelled_by_previous_deploy(self):
        from app.services.retired_work_point import purge_retired_library_examples
        from app.services.remark_catalog import normalize
        old_text='Балкон-очист.от раствора'
        digest=hashlib.sha256(normalize(old_text).encode()).hexdigest()
        entry=RemarkLibraryEntry(normalized_hash=digest,point_key='25',text=old_text,
                                 points_json='["25"]',occurrences=1)
        db.session.add(entry)
        AppSetting.query.filter_by(key='retired_point_library_cleanup_v2').delete()
        db.session.commit();purge_retired_library_examples()
        self.assertIsNone(RemarkLibraryEntry.query.filter_by(normalized_hash=digest).first())
        path=Path(self.directory.name)/'old-relabeled.json'
        path.write_text(json.dumps({'schema':SCHEMA,'entries':[{'text':old_text,'points':['25']}]}))
        import_library(path)
        self.assertIsNone(RemarkLibraryEntry.query.filter_by(normalized_hash=digest).first())

    def test_real_startup_migrates_existing_database_once(self):
        root=Path(self.directory.name)
        class FileConfig(Config):
            TESTING=True
            SECRET_KEY='startup-retirement'
            SQLALCHEMY_DATABASE_URI='sqlite:///'+str(root/'old.sqlite')
            UPLOAD_FOLDER=str(root/'startup-uploads')
            EXPORT_FOLDER=str(root/'startup-exports')
        app=create_app(FileConfig)
        with app.app_context():
            project=Project(name='Existing server project')
            point=WorkPoint(point_number='16',short_name='Разнорабочие')
            task=Task(project=project,work_point=point,source_uid='existing-server',description='Контрольный неизвестный дефект')
            db.session.add_all([project,point,task])
            AppSetting.query.filter_by(key=MARKER).delete()
            db.session.commit(); task_id=task.id;db.session.remove()
        for _ in range(2):
            restarted=create_app(FileConfig)
            with restarted.app_context():
                self.assertFalse(WorkPoint.query.filter_by(point_number='16').one().is_active)
                self.assertIsNotNone(db.session.get(Task,task_id))
                self.assertEqual(Task.query.count(),1)
                self.assertTrue(db.session.get(Task,task_id).is_archived)
                self.assertFalse(db.session.info.get('retiring_work_point'))
                db.session.remove()

    def test_actual_download_routes_return_clean_xlsx_and_forms_omit_point(self):
        from io import BytesIO
        from app import login_manager
        path=self.source()
        self.task('route-pvc',self.pvc,'Регулировка ПВХ')
        self.migrate()
        self.user.all_projects_access=True
        set_setting(f'latest_excel_path_project_{self.project.id}',str(path))
        set_setting('latest_excel_path',str(path));db.session.commit()
        client=self.app.test_client()
        old_protection=login_manager.session_protection
        login_manager.session_protection=None
        try:
            with client.session_transaction() as session:
                session.update(_user_id=self.user.get_auth_id(),_fresh=True,
                               session_version=self.user.session_version or 0,current_project_id=self.project.id)
            for url in [f'/apartments/{self.apartment.id}/source-table/export','/export/source-with-strikes']:
                with self.subTest(url=url):
                    response=client.get(url)
                    self.assertEqual(response.status_code,200)
                    self.assertIn('attachment',response.headers['Content-Disposition'])
                    workbook=load_workbook(BytesIO(response.data),data_only=True)
                    self.assert_no_retired_header(workbook)
                    workbook.close();response.close()
            response=client.get('/tasks/new?mode=manual&kind=single')
            self.assertEqual(response.status_code,200)
            self.assertNotIn('<option value="16"',response.get_data(as_text=True))
            self.assertIn('<option value="17"',response.get_data(as_text=True))
            response=client.post('/tasks/new',data={'add_mode':'manual','manual_kind':'single',
                'apartment_id':str(self.apartment.id),'point_number':'16','description':'Устаревший пункт'})
            self.assertIn(response.status_code,(200,302))
            self.assertIsNone(Task.query.filter_by(description='Устаревший пункт').first())
        finally:
            login_manager.session_protection=old_protection

    def test_full_export_shifts_formulas_filters_and_merged_headers(self):
        from app.services.excel_export import _remove_retired_columns
        workbook=Workbook();sheet=workbook.active;sheet.title='Таблица'
        sheet.append(['Квартира','Разнорабочие','ПВХ','Итого'])
        sheet.append([42,99,5,'=SUM(B2:C2)'])
        sheet['D3']='=C2+$C$2+COUNTIF(B:B,">0")'
        sheet.auto_filter.ref='A1:D3'
        sheet.merge_cells('C4:D4');sheet['C4']='Сохранить объединение'
        other=workbook.create_sheet('Другое')
        other['A1']="='Таблица'!$C$2"
        other['A2']='=B2';other['A3']='="C2 is literal"'
        _remove_retired_columns(workbook)
        self.assertEqual(sheet['C2'].value,'=SUM(B2:B2)')
        self.assertEqual(sheet['C3'].value,'=B2+$B$2+COUNTIF(0,">0")')
        self.assertEqual(other['A1'].value,"='Таблица'!$B$2")
        self.assertEqual(other['A2'].value,'=B2')
        self.assertEqual(other['A3'].value,'="C2 is literal"')
        self.assertEqual(sheet.auto_filter.ref,'A1:C3')
        self.assertIn('B4:C4',{str(region) for region in sheet.merged_cells.ranges})
        self.assertEqual(sheet['B4'].value,'Сохранить объединение')

    def test_legacy_column_sixteen_with_pvc_name_is_preserved_not_retired(self):
        legacy=WorkPoint(point_number='16',short_name=WORK_POINT_LABELS['17'],source_sheet_name='legacy')
        task=self.task('real-pvc-legacy',legacy,'Повреждение ПВХ')
        task_id=task.id; source_uid=task.source_uid
        result=self.migrate();db.session.expire_all()
        kept=db.session.get(Task,task_id)
        self.assertEqual(kept.work_point.point_number,'17')
        self.assertEqual(kept.source_uid,source_uid)
        self.assertFalse(kept.is_missing_in_latest_sync)
        self.assertEqual(result['moved'],0)
        self.assertEqual(Task.query.count(),1)
